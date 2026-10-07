from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
SEARCH=ROOT/'results/20261006-093705-631888'
FINAL=ROOT/'results_best_rates/20261006-113427-959496'
c=json.loads((SEARCH/'config.json').read_text())
f=json.loads((FINAL/'config.json').read_text())
runs=pd.read_csv(SEARCH/'runs.csv')
summary=pd.read_csv(SEARCH/'summary.csv').set_index('method')
final=pd.read_csv(FINAL/'final_accuracy.csv').set_index('method')
assert len(runs)==480 and (runs.status=='ok').all()
history=pd.read_csv(SEARCH/'history.csv')
last=history[history.epoch==40]
assert last.groupby(['method','base_lr']).seed.nunique().eq(20).all()
score=last.groupby(['method','base_lr']).val_accuracy.mean()
labels={'sgd':'SGD','adam':'Adam','inverse_mag':'Inverse MAG','adagrad_norm':'AdaGrad Norm','mag':'MAG','mag_floor':'MAG floor'}
order=['sgd','adam','inverse_mag','adagrad_norm','mag','mag_floor']
rates={v['method']:v['base_lr'] for v in f['selected']}
for m in order:
 assert rates[m]==score.loc[m].idxmax()
D=Document()
s=D.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=D.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(10.5)
D.styles['Normal'].paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.08
D.styles['Title'].font.size=Pt(25)
D.styles['Heading 1'].font.size=Pt(18)
D.styles['Heading 2'].font.size=Pt(12)
D.core_properties.author='Charis Aletraris'
D.core_properties.title='Gradient magnitude optimizer experiment on Iris'
foot=s.footer.paragraphs[0];foot.alignment=2
r=foot.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
def p(t): return D.add_paragraph(t)
def h(t): D.add_heading(t,1)
def sub(t): D.add_heading(t,2)
def page(t): D.add_page_break();h(t)
def eq(t):
 pp=D.add_paragraph();pp.alignment=1
 om=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath');rr=OxmlElement('m:r');tt=OxmlElement('m:t');tt.text=t;rr.append(tt);math.append(rr);om.append(math);pp._p.append(om)
def table(headers, rows, widths):
 t=D.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for cell,w in zip(t.columns,widths):cell.width=Inches(w)
 for cell,txt in zip(t.rows[0].cells,headers):cell.text=txt
 for row in rows:
  for cell,txt in zip(t.add_row().cells,row):cell.text=str(txt)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if i==0 else ('F4F6F8' if i%2==0 else 'FFFFFF'));pr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   pr.append(borders)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_before=Pt(3);pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.line_spacing=1
    pp.alignment=0 if j==0 else 1
    for r in pp.runs:r.font.size=Pt(9);r.bold=(i==0)
 return t

def picture(name,caption):
 D.add_picture(str(FINAL/name),width=Inches(6.8));D.add_paragraph(caption,'Caption')

D.add_paragraph('Gradient magnitude optimizer experiment on Iris','Title')
p('Charis Aletraris | 6 October 2026')
p('This report compares six optimizers using the same Iris neural network. It documents the learning-rate search, the validation-based choice of one base rate per optimizer, and the resulting training, validation and test accuracy curves. Base rates are now selected by the highest mean validation accuracy at epoch 40 across 20 runs. With this criterion, my original MAG rule achieved 98.17% mean test accuracy and Adam achieved 97.67% after 40 epochs.')
h('Algorithms and learning rate values')
p('The base rate is a constant hyperparameter. For adaptive rules, the effective learning rate can change at every minibatch. Four candidates per optimizer were tested; each candidate was trained with 20 seeds.')
table(['Optimizer','Tested base rates','Selected'],[[labels[m],', '.join(f'{v:g}' for v in c['grids'][m]),f'{rates[m]:g}'] for m in order],[1.3,4.3,1.2])
sub('Reference methods')
p('SGD uses a fixed learning rate and zero momentum. Adam uses TensorFlow defaults apart from its base rate: beta 1 = 0.9, beta 2 = 0.999 and epsilon = 0.0000001; it adapts updates through first and second gradient moments.')
p('Inverse MAG uses the professorâ€™s suggested denominator variant: eta_t = base_rate / (m_t + epsilon), with epsilon = 0.00000001. AdaGrad Norm accumulates the squared global gradient norm, starts with b_0 = 1, and applies eta_t = base_rate / b_t after adding the current gradient norm to b_t squared.')
p('RMSProp and coordinate-wise AdaGrad are available in the code but were not included in these six-method experiments. LARS and Polyak-type step sizes were discussed as related approaches and were not implemented.')
sub('My two proposed rules')
p('MAG multiplies the base rate by the mean absolute gradient. MAG floor uses the same multiplier when it exceeds 1 and uses 1 otherwise. Their exact equations are given next.')

page('Equations for my two algorithms')
p('For minibatch t, L_t is mean categorical cross-entropy and theta_t contains all trainable weights and biases. The gradient is taken after averaging the minibatch loss, rather than averaging absolute per-example gradients.')
eq('gâ‚œ = âˆ‡Lâ‚œ(Î¸â‚œ)       mâ‚œ = (1/N) âˆ‘áµ¢â‚Œâ‚á´º |gâ‚œ,áµ¢|')
p('The network has N = (4 Ã— 16 + 16) + (16 Ã— 3 + 3) = 131 trainable scalars. Thus m_t is the sum of the absolute values of all 131 gradient components divided by 131, across all layers.')
sub('MAG with the original multiplier')
eq('Î·â‚œ = Î·â€² mâ‚œ       Î¸â‚œâ‚Šâ‚ = Î¸â‚œ âˆ’ Î·â‚œ gâ‚œ       Î·â€² = 10')
p('Large gradient magnitudes increase both the gradient and its multiplier; small gradient magnitudes reduce both. The implementation does not add clipping, smoothing, momentum or a learning-rate cap to this rule.')
sub('MAG floor with a minimum multiplier of one')
eq('mÌƒâ‚œ = max(mâ‚œ, 1)       Î·â‚œ = Î·â€² mÌƒâ‚œ       Î¸â‚œâ‚Šâ‚ = Î¸â‚œ âˆ’ Î·â‚œ gâ‚œ       Î·â€² = 1')
p('If m_t is greater than 1, the original MAG multiplier is used. If m_t is at most 1, the multiplier is 1 and the update equals plain SGD with the same base rate. A zero gradient still produces a zero update.')
h('Experimental setup')
p('I used the 150-sample Iris dataset with four input features and three classes. The network has 4 inputs, one hidden layer of 16 tanh units, and 3 softmax outputs. Labels are one-hot encoded. Training uses mean categorical cross-entropy, batch size 32 and 40 epochs.')
p('The assignment split is reproduced exactly: 108 training samples, 12 validation samples and 30 test samples. A non-stratified 80/20 split with random_state = 42 produces 120 non-test rows; the last 12 are reserved for validation before shuffling. StandardScaler is fitted on all 120 non-test rows, as in the assignment. Validation features therefore contribute to scaling statistics, although test features do not.')
p('Seeds 42â€“61 provide 20 independent initializations and batch orders on one fixed data split. For a given seed, all methods and candidates start with the same weights and use the same minibatch order. TensorFlow 2.21.0 was used. The search contains 6 Ã— 4 Ã— 20 = 480 runs; the fixed-rate comparison contains 6 Ã— 20 = 120 runs.')

page('Learning rate search results')
p('For each run_experiment.py candidate and seed, validation accuracy at epoch 40 was extracted from history.csv. The table reports the mean final validation accuracy across all 20 seeds. Higher is better. This replaces the earlier selection based on the mean best validation loss reached at any epoch. Every one of the 480 search runs completed with status ok; this indicates numerical completion, not necessarily good accuracy.')
table(['Optimizer','Base rate','Mean final validation accuracy','Chosen'],[[labels[m],f'{rate:g}',f'{100*score.loc[m,rate]:.2f}%','Yes' if rate==rates[m] else ''] for m in order for rate in c['grids'][m]],[1.55,1.0,3.15,1.1])
p('The selected value is the highest-scoring tested candidate for each optimizer. MAG and inverse MAG select the upper boundary of their grids, so larger or more finely spaced rates could still be worth investigating. These are the best tested values, not proven global optima.')

page('Selection rationale and numerical results')
p('One base rate per optimizer was selected by averaging validation accuracy at epoch 40 over all 20 seeds and choosing the largest average. A tie would favor the smaller rate; no winning ties occurred. This criterion matches the aim of comparing accuracy after a fixed 40-epoch budget. Test accuracy was not used for rate selection. Adam changes from 0.1 to 0.01, and AdaGrad Norm changes from 10 to 1. The other four rates remain unchanged.')
p('The original search summary below is retained as a historical comparison: it selected both the rate and best epoch separately for each seed using validation loss. Its test summary therefore describes per-seed selected checkpoints. The updated runner uses final mean validation accuracy to fix one rate per optimizer and reports epoch 40 without restoring an earlier checkpoint. These two result tables measure different selection procedures.')
sub('Earlier loss selected checkpoint results')
table(['Optimizer','Mean test accuracy','SD across seeds'],[[labels[m],f'{100*summary.loc[m,"test_accuracy_mean"]:.2f}%',f'{100*summary.loc[m,"test_accuracy_std"]:.2f} pp'] for m in order],[2.4,2.2,2.2])
sub('Fixed rate results after 40 epochs')
table(['Optimizer','Base rate','Training','Validation','Test'],[[labels[m],f'{rates[m]:g}',*[f'{100*final.loc[m,k]:.2f}%' for k in ['train_accuracy','val_accuracy','test_accuracy']]] for m in order],[1.65,.85,1.4,1.45,1.45])
p('All accuracies are arithmetic means across 20 runs. In the fixed-rate results, MAG has the highest mean test accuracy at 98.17%; Adam follows at 97.67%, a difference of 0.50 percentage points. Inverse MAG and AdaGrad Norm each reach 97.50%, while SGD and MAG floor each reach 96.67%. MAG floor and SGD also match in final training and validation accuracy.')
p('The mean test score aggregates repeated predictions on the same 30 test samples. It is not evidence from 600 independent test cases. The following figures show every epoch, rather than only the final values.')

page('Training and validation accuracy curves')
p('Each curve uses the fixed base rate in the legend and averages all 20 seeds at each epoch. Training accuracy is evaluated on the full training partition after the epochâ€™s updates.')
picture('training_accuracy.png','Figure 1  Mean training accuracy at each epoch with one selected base rate per optimizer')
picture('validation_accuracy.png','Figure 2  Mean validation accuracy at each epoch on the fixed 12-sample validation partition')

page('Test accuracy and interpretation')
picture('test_accuracy.png','Figure 3  Mean test accuracy at each epoch on the fixed 30-sample test partition')
p('The test set is evaluated after every epoch for reporting only. These evaluations do not affect gradients, stopping, checkpoint restoration or learning-rate selection. The fixed-rate rerun reuses the original seeds and split, so it reproduces the selected setting rather than providing an independent confirmation study.')
sub('What the results support')
p('My original MAG rule has the highest mean final test accuracy in this comparison, 0.50 percentage points above Adam. This small difference does not establish a consistent or statistically significant advantage. MAG floor matches SGDâ€™s final metrics, consistent with its SGD-like behavior whenever the mean absolute gradient is at most 1. Matching final accuracies alone does not establish that every intermediate update was identical.')
p('The findings concern one small dataset, architecture and fixed partition. The validation set has only 12 samples and influences preprocessing statistics. The learning-rate search was expanded after earlier results were inspected, so the test set is not a fresh final holdout for the overall exploratory process. Further comparisons should use training-only preprocessing, additional datasets and splits, fixed search budgets and an untouched final test set before claiming general superiority or novelty.')
sub('Sources and reproducibility')
p('Learning-rate search: results/20261006-093705-631888, using config.json, history.csv, runs.csv and summary.csv. Fixed-rate comparison: results_best_rates/20261006-113427-959496, using config.json, final_accuracy.csv and the three original PNG plots. All paths are relative to gradient-magnitude-experiment.')
p('Implementation: optimizers.py, run_experiment.py and run_best_rates.py. The selected rates are now written directly in BEST_BASE_RATES in run_best_rates.py. The JSON selection record uses the same rates. Related reference: Ward, R., Wu, X. and Bottou, L. (2019), AdaGrad Stepsizes: Sharp Convergence Over Nonconvex Landscapes, PMLR 97, 6677â€“6686. https://proceedings.mlr.press/v97/ward19a.html')
output=ROOT/'Gradient_Magnitude_Experiment_Report_Updated.docx'
D.save(output)
print(output)

