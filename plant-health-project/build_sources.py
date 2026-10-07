from pathlib import Path
import tarfile, json, hashlib, csv
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

root = Path(__file__).resolve().parent
flowers = root/'datasets/oxford-flowers102'
matdata = json.loads((flowers/'labels_intermediate.json').read_text())
labels = matdata['imagelabels.mat']['labels']
splitdata = matdata['setid.mat']
splits = {int(i): s for k,s in [('trnid','train'),('valid','validation'),('tstid','test')] for i in splitdata[k]}
with tarfile.open(flowers/'102flowers.tgz','r:gz') as archive:
    members = [m for m in archive if m.isfile() and m.name.endswith('.jpg')]
    for m in members:
        f = archive.extractfile(m)
        while f.read(1024*1024): pass
assert len(members) == len(labels) == 8189
with (flowers/'image_labels.csv').open('w',newline='',encoding='utf-8') as f:
    w=csv.writer(f); w.writerow(['image_path_in_archive','class_id','split'])
    for i,l in enumerate(labels,1): w.writerow([f'jpg/image_{i:05}.jpg',int(l),splits[i]])
(flowers/'verification.json').write_text(json.dumps({'images':8189,'classes':len(set(labels)),'archive_read_check':'passed','sha256':hashlib.file_digest((flowers/'102flowers.tgz').open('rb'),'sha256').hexdigest()},indent=2))

doc=Document()
sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Cm(2); sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2']:
    doc.styles[name].font.name='Arial'; doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(15)
def p(t): doc.add_paragraph(t)
def link(label,url):
    para=doc.add_paragraph(); h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),para.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r); para._p.append(h)
    # Keep the exact URL in the document for unambiguous provenance.
    r=para.add_run('\n'+url); r.font.size=Pt(8)

doc.add_paragraph('Πηγές δεδομένων φυτών',style='Title')
p('Deep Learning • Καταγραφή λήψεων 27 Σεπτεμβρίου 2026')
p('Τα δεδομένα αποθηκεύτηκαν στον φάκελο deep-learning-course/plant-health-project. Το PlantVillage υποστηρίζει αναγνώριση καλλιέργειας και κατάστασης φύλλου. Το Oxford Flowers 102 υποστηρίζει αναγνώριση κατηγορίας άνθους. Κανένα από τα δύο δεν παρέχει τεκμηριωμένη θεραπεία ανά φωτογραφία.')
doc.add_heading('PlantVillage',1)
p('Αρχείο: datasets/plantvillage/data.zip — 2.184.723.441 bytes. Η λήψη ολοκληρώθηκε και ο έλεγχος CRC όλων των αρχείων ZIP πέρασε. Η έγχρωμη έκδοση περιέχει 54.305 εικόνες και 38 κατηγορίες. Το ZIP περιέχει επίσης ασπρόμαυρες και απομονωμένες εκδοχές φύλλων, οι οποίες δεν είναι ανεξάρτητα νέα δείγματα.')
p('Οι ετικέτες βρίσκονται στα ονόματα φακέλων, π.χ. Raspberry___healthy. Δημιουργήθηκαν τα image_labels.csv και class_counts.csv και 38 ενδεικτικές φωτογραφίες στον φάκελο examples/plantvillage. Τα αρχεία CSV είναι τοπικά παράγωγα των ονομάτων μέσα στο ZIP.')
link('Πηγή και δημιουργοί — αποθετήριο spMohanty', 'https://github.com/spMohanty/PlantVillage-Dataset')
link('Ακριβής σύνδεσμος λήψης — Hugging Face του Mohanty','https://huggingface.co/datasets/mohanty/PlantVillage/resolve/main/data.zip')
p('Αναφορά: Mohanty, Hughes και Salathé (2016), Using Deep Learning for Image-Based Plant Disease Detection. DOI: 10.3389/fpls.2016.01419. Διατηρήστε την αναφορά και ελέγξτε τους όρους της πηγής πριν από αναδιανομή. Για train/test, οι φωτογραφίες του ίδιου φύλλου και οι παραλλαγές τους πρέπει να παραμένουν στην ίδια ομάδα.')
doc.add_heading('Oxford Flowers 102',1)
p('Αρχείο: datasets/oxford-flowers102/102flowers.tgz — 344.862.509 bytes. Περιέχει 8.189 εικόνες σε 102 κατηγορίες ανθών. Η ανάγνωση του πλήρους συμπιεσμένου αρχείου και η αντιστοίχιση με τις ετικέτες ελέγχθηκαν. Δεν είναι γενικό dataset όλων των ειδών φυτών ούτε dataset ασθενειών.')
p('Κατέβηκαν επίσης imagelabels.mat, setid.mat και README.txt. Το τοπικό image_labels.csv συνδέει κάθε φωτογραφία με το αριθμητικό class_id και το επίσημο train/validation/test split.')
link('Πρωτογενής πηγή — Visual Geometry Group University of Oxford','https://www.robots.ox.ac.uk/~vgg/data/flowers/102/')
link('Ακριβής σύνδεσμος λήψης εικόνων','https://www.robots.ox.ac.uk/~vgg/data/flowers/102/102flowers.tgz')
p('Τα συνοδευτικά αρχεία προέρχονται από την ίδια διεύθυνση, με το αντίστοιχο όνομα αρχείου αντί του 102flowers.tgz. Αναφορά: Nilsback και Zisserman (2008), Automated Flower Classification over a Large Number of Classes. Οι όροι χρήσης ελέγχονται στην επίσημη πηγή πριν από αναδιανομή.')

doc.add_page_break()
doc.add_heading('Πρόσθετες πηγές για επέκταση',1)
p('Οι ακόλουθες πηγές καταγράφηκαν για μελλοντική χρήση. Τα πλήρη datasets τους δεν έχουν κατεβεί.')
doc.add_heading('Flavia για αναγνώριση φύλλων',2)
p('Η επίσημη πηγή περιγράφει 32 είδη και πίνακα αντιστοίχισης ονομάτων με αριθμούς φωτογραφιών. Αποθηκεύτηκε το source-labels.html. Οι δύο σύνδεσμοι λήψης επέστρεψαν HTML αντί του αρχείου δεδομένων. Το υπάρχον Leaves.tar.bz2 δεν είναι έγκυρο dataset και δεν πρέπει να χρησιμοποιηθεί. Παραμένει στον φάκελο σύμφωνα με την οδηγία να μη διαγραφεί κανένα αρχείο.')
link('Επίσημη πηγή και ετικέτες','https://flavia.sourceforge.net/')
link('Αρχικός σύνδεσμος λήψης που δοκιμάστηκε','https://downloads.sourceforge.net/project/flavia/Leaf%20Image%20Dataset/1.0/Leaves.tar.bz2')
link('Εναλλακτικός σύνδεσμος που δοκιμάστηκε','https://netix.dl.sourceforge.net/project/flavia/Leaf%20Image%20Dataset/1.0/Leaves.tar.bz2')
doc.add_heading('Ασθένειες και υδατική καταπόνηση',2)
p('PlantDoc: 2.598 δείγματα, 13 είδη φυτών και έως 17 κατηγορίες ασθενειών, από φωτογραφίες εκτός εργαστηρίου που συγκεντρώθηκαν από το διαδίκτυο. Άδεια CC BY 4.0. Χρήσιμο για αξιολόγηση με πιο σύνθετο φόντο· δεν παρέχει θεραπεία ανά εικόνα.')
link('PlantDoc — αποθετήριο των δημιουργών','https://github.com/pratikkayal/PlantDoc-Dataset')
p('Maize leaf dataset: 18.040 αποκόμματα από 656 αρχικές φωτογραφίες, σε τρία επίπεδα ποτίσματος. Άδεια CC BY 4.0. Οι ετικέτες περιγράφουν πειραματικές συνθήκες, όχι ακριβή συνταγή ποτίσματος. Τα αποκόμματα της ίδιας αρχικής εικόνας πρέπει να μένουν στο ίδιο split.')
link('Shuo Zhuang — Mendeley Data 2019','https://data.mendeley.com/datasets/w77b87fvct/1')
doc.add_heading('Γενική αναγνώριση και κατηγορία φυτού',2)
p('Pl@ntNet 300K v2: 306.087 εικόνες και 1.000 είδη, με ετικέτες οργάνου όπως φύλλο, άνθος και καρπός. Το αρχείο εικόνων είναι 41,8 GB και δεν κατέβηκε. Η άδεια διαφέρει ανά εικόνα. Είναι πιο κοντά στις φυσικές φωτογραφίες χρηστών, αλλά το τεκμηριωμένο σχήμα δεν περιέχει κατηγορία δέντρο/θάμνος/ποώδες.')
link('Pl@ntNet 300K v2 — επίσημη εγγραφή Zenodo','https://zenodo.org/records/10419064')
p('Για δέντρο, θάμνο ή ποώδες χρειάζεται πρόσθετη, τεκμηριωμένη αντιστοίχιση ανά επιστημονικό όνομα. Δεν προστέθηκαν αυθαίρετες ετικέτες. Το άνθος είναι όργανο φυτού και δεν αποτελεί αμοιβαία αποκλειόμενη κατηγορία με το δέντρο ή τον θάμνο. Ομοίως, η ανάγκη για περισσότερο νερό ή ήλιο δεν προκύπτει αυτόματα από την ετικέτα ασθένειας.')
doc.save(root/'Dataset_Sources.docx')
print('Created Dataset_Sources.docx; Oxford archive and labels verified.')
