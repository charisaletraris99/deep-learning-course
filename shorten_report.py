from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).parent
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.3), Inches(11.7)
s.top_margin = s.bottom_margin = Inches(0.6)
s.left_margin = s.right_margin = Inches(0.75)
for name in ['Normal', 'Title', 'Heading 1', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(5)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1
doc.styles['Title'].font.size = Pt(20)
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(7)

doc.add_paragraph('Exercise 1a Learning rate', 'Title')
doc.add_heading('Model and training', 1)
doc.add_paragraph('We used the Iris dataset with four inputs, one hidden layer of 16 neurons using tanh, and three outputs using softmax. We trained with SGD, momentum 0, categorical cross-entropy loss and batch size 32. The data contained 108 flowers for training, 12 for validation and 30 for testing. Labels were one-hot encoded and inputs were standardized using the training partition.')
doc.add_paragraph('We used 40 epochs and 20 fresh runs per learning rate, as required by the assignment. An epoch is one pass through the training data. The fixed 40 epochs give each rate the same training budget. Averaging 20 runs reduces the effect of random starting weights. Each plotted point is the average test accuracy at that epoch.')
doc.add_heading('Mean accuracy after 40 epochs', 1)
rows = [('Learning rate', 'Training', 'Validation', 'Test'), ('0.003', '73.66%', '68.75%', '76.67%'), ('0.03', '88.33%', '88.75%', '94.33%'), ('0.06', '92.82%', '91.67%', '96.00%'), ('0.2', '96.16%', '95.00%', '98.00%'), ('0.5', '96.30%', '96.25%', '97.67%')]
t = doc.add_table(rows=0, cols=4)
for i, values in enumerate(rows):
    for cell, value in zip(t.add_row().cells, values):
        cell.text = value
        props = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for side in ['top', 'left', 'bottom', 'right']:
            b = OxmlElement('w:'+side)
            for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: b.set(qn('w:'+k),v)
            borders.append(b)
        props.append(borders)
        if i == 0:
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EDF2'); props.append(shade)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs: r.bold = i == 0
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.keep_with_next = True
p.add_run().add_picture(str(root/'exercise1a_test_accuracy.png'), width=Inches(4.65))
p = doc.add_paragraph('Mean test accuracy over 20 runs for each learning rate.', 'Caption')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('Which learning rate performed best', 1)
doc.add_paragraph('The best rate was 0.2, with 98.00% mean test accuracy at epoch 40. The rate of 0.5 reached 97.67%, only 0.33 percentage points lower. Their final performance was therefore very similar.')
doc.add_heading('What the learning rate does', 1)
doc.add_paragraph('The learning rate controls how much the weights change after each batch. Higher rates generally improved training, validation and test accuracy faster in this experiment. A low rate takes smaller steps and may need more epochs to reach high accuracy.')
doc.add_paragraph('At 0.5, the larger changes can overshoot good weight values or cause fluctuations. This may explain its slightly lower final test accuracy than 0.2, but the small difference does not prove the cause. Extra epochs do not always stabilize a high rate; lowering the rate may help. Here, the 0.5 curve already levels off early.')
doc.save(root/'Exercise_1a_Learning_Rate_Report.docx')
print('Short report saved: 1 embedded plot and 1 results table.')
