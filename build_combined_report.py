from pathlib import Path
import re
import json
from pypdf import PdfReader
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).parent
reader=PdfReader(r'C:/Users/Charis Local/Downloads/Assignment1_Charis_Aletraris_report.docx.pdf')
texts=[re.sub(r'\s+', ' ', page.extract_text()).strip() for page in reader.pages]
assets=root/'report_render'
assets.mkdir(exist_ok=True)
for i, im in enumerate(reader.pages[1].images):
    im.image.save(assets/f'original_plot_{i}.png')

doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(0.65)
s.left_margin=s.right_margin=Inches(0.75)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name]
    st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.06
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 2'].font.size=Pt(12)
def p(txt): doc.add_paragraph(txt)
def h(txt): doc.add_heading(txt,level=2)
def fragment(page,start,end=None):
    t=texts[page]; a=t.index(start); b=t.index(end,a) if end else len(t)
    return t[a:b].strip()
def table(rows):
    t=doc.add_table(rows=0,cols=4)
    for i,row in enumerate(rows):
        for c,value in zip(t.add_row().cells,row):
            c.text=str(value)
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side)
                for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: b.set(qn('w:'+k),v)
                borders.append(b)
            pr.append(borders)
            if i==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EDF2'); pr.append(sh)
            for para in c.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_before=Pt(3)
                para.paragraph_format.space_after=Pt(3)
                for run in para.runs: run.bold=(i==0)
def picture(path,width):
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path),width=Inches(width))
def pair(paths):
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for path in paths: para.add_run().add_picture(str(path),width=Inches(3.35))

doc.add_paragraph('Exercise 1',style='Title')
h('Model and experimental procedure')
for start,end in [('The model takes','We used stochastic'),('We used stochastic','The 150 flowers'),('The 150 flowers','Why 40 epochs')]: p(fragment(0,start,end))
h('Why 40 epochs and 20 runs')
p(fragment(0,'Both settings','Each rate was'))
p(fragment(0,'Each rate was','A) Learning rate'))
doc.add_heading('A Learning rate',level=1)
h('Mean accuracy after 40 epochs')
table([['Learning rate','Training','Validation','Test'],['0.003','73.66%','68.75%','76.67%'],['0.03','88.33%','88.75%','94.33%'],['0.06','92.82%','91.67%','96.00%'],['0.2','96.16%','95.00%','98.00%'],['0.5','96.30%','96.25%','97.67%']])
doc.add_page_break()
h('Training and validation accuracy curves')
pair([assets/'original_plot_0.png',assets/'original_plot_1.png'])
p(fragment(1,'As we can see','a)Which learning'))
picture(assets/'original_plot_2.png',4.65)
h('a Which learning rate performed best')
p(fragment(1,'Using mean test accuracy','The saved curves'))
p(fragment(1,'The saved curves'))
doc.add_page_break()
h('b What the learning rate is and how it affects training')
for start,end in [('The learning rate controls','Small updates'),('Small updates','A plausible'),('A plausible','A high rate'),('A high rate',None)]: p(fragment(2,start,end))

doc.add_page_break()
doc.add_heading('B Momentum',level=1)
h('Model and experimental procedure')
p('We used the same network and data preparation as in Exercise 1(a): four inputs, 16 hidden neurons with tanh, and three outputs with softmax. The learning rate was fixed at 0.02, and we compared momentum values of 0, 0.3, 0.6, 0.9 and 0.99. Batch size remained 32, with categorical cross-entropy loss and a validation split of 0.1.')
p('Each momentum value was trained for 40 epochs in 20 fresh runs, using seeds 42 to 61 and the same data split. The fixed number of epochs gives each value the same training budget. Averaging the runs reduces the effect of random starting weights. Each curve shows the mean accuracy at every epoch.')
data_file=root/'report_render'/'momentum_summary.json'
if data_file.exists():
    h('Mean accuracy after 40 epochs')
    table([['Momentum','Training','Validation','Test']]+json.loads(data_file.read_text()))
h('Training and validation accuracy curves')
pair([root/'exercise1b_training_accuracy.png',root/'exercise1b_validation_accuracy.png'])
p('Higher momentum generally helped the model reach high accuracy sooner in this experiment. Values of 0.9 and 0.99 learned faster than 0, 0.3 and 0.6. The training accuracy for 0.99 ended highest, but its validation curve fluctuated and fell after an earlier peak. By epoch 40, 0.9 had the highest validation accuracy. Higher momentum therefore did not improve every measure in the same way.')
doc.add_page_break()
h('Test accuracy curves')
picture(root/'exercise1b_test_accuracy.png',5.6)
h('a Which momentum performed best')
p('Using mean test accuracy at epoch 40, momentum 0.99 performed best. Its test curve finished just below 100%, above the 0.9 curve at about 98%. These percentages are approximate readings from the plot. Momentum 0.9 reached high test accuracy earlier and followed a smoother curve, but 0.99 finished with the highest value. This is the best result for the tested setup, rather than a guarantee that 0.99 will always be best.')
h('b What momentum is and how it affects training')
p('Momentum keeps part of the previous weight update and adds it to the update calculated from the current batch. It helps the model keep moving in a useful direction, like a ball building speed when pushed repeatedly in the same direction. A value of 0 uses no previous update, while 0.9 carries 90% of the previous update into the next one.')
p('Momentum acts between batches and continues across epochs. It resets between the 20 independent runs because each run creates a new model and optimizer. It does not carry over the learning rate or the previous run\'s accuracy.')
p('A larger momentum can speed up learning and reduce some back-and-forth movement. However, if it is too strong, the model can overshoot useful weight values and fluctuate. The fluctuations at 0.99 are consistent with this effect, although the curves alone do not prove the cause. In this experiment, 0.99 gave the highest final test accuracy, while 0.9 offered smoother progress and the highest final validation accuracy.')
out=root/'Assignment1_Charis_Aletraris_Report_1a_1b.docx'
doc.save(out)
print(out)
print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables),'Images:',len(doc.inline_shapes))
