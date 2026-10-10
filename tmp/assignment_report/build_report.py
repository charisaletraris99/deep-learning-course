from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root=Path.cwd()
plots=root/'assignment-2/outputs/exercise_1_three_conv_gelu'
resnet_plots=root/'assignment-2/outputs/exercise_2_residual_gelu'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    style=doc.styles[name]
    style.font.name='Calibri'; style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.08
for name,size in [('Title',23),('Heading 1',16),('Heading 2',12)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(10)
doc.styles['Caption'].font.size=Pt(9)
doc.core_properties.title='Fashion MNIST CNN and ResNet results'
doc.core_properties.subject='Assignment 2 Exercises 1 and 2'
def p(text): return doc.add_paragraph(text)
def h(text): doc.add_heading(text,level=1)
def point(label,text):
    para=doc.add_paragraph(); para.add_run(label+' ').bold=True; para.add_run(text)
def fig(name,width,caption,source=plots):
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    run=para.add_run(); run.add_picture(str(source/name),width=Inches(width))
    para=doc.add_paragraph(caption,style='Caption')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for cell,text in zip(t.rows[0].cells,headers): cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            pr.append(borders)
            margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side); v.set(qn('w:w'),'85'); v.set(qn('w:type'),'dxa'); margin.append(v)
            pr.append(margin)
            if i==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EDF3'); pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                for run in para.runs: run.font.size=Pt(10); run.bold=(i==0)
    return t

doc.add_paragraph('Fashion MNIST CNN and ResNet results',style='Title')
p('Assignment 2  |  Exercises 1 and 2')
p('We trained a convolutional neural network to classify ten types of clothing. This report explains our model choices and the results. The loss curves show signs of overfitting from around epoch 6, even though validation and test accuracy remain fairly steady.')
h('Data and model structure')
p('Each input is a 28 x 28 grayscale image. We use 54,000 images for training, 6,000 for validation and 10,000 for testing. Validation helps us choose the model weights. The test set measures performance on images not used to update the weights.')
p('We divide pixel values by 255 to put them between 0 and 1. We one-hot encode each label into ten values and add one grayscale channel before the first convolution. Seed 42 makes the data split repeatable.')
table(['Stage','Main settings','Output per image'],[
('Convolution 1','32 filters, 3 x 3, stride 1, GELU','28 x 28 x 32'),
('Max pooling 1','2 x 2, stride 2','14 x 14 x 32'),
('Convolution 2','64 filters, 3 x 3, stride 1, GELU','14 x 14 x 64'),
('Max pooling 2','2 x 2, stride 2','7 x 7 x 64'),
('Convolution 3','64 filters, 3 x 3, stride 1, GELU','7 x 7 x 64'),
('Classifier','Flatten, Dense 128, GELU, dropout 0.3','128 values'),
('Output','Dense 10, softmax','10 probabilities')],[1.35,3.5,1.95])
p('All three convolutions use one-pixel replicate padding. The complete model has 458,570 trainable parameters.')
h('Why we use these convolutions')
point('Three layers and 3 x 3 kernels.', 'We want to learn local patterns first and then combine them in deeper layers. A 3 x 3 kernel examines a position and its immediate neighbors. Deeper layers combine information from a larger area of the original image.')
point('Stride 1.', 'The filter moves one position at a time. We avoid reducing the image size during convolution and leave spatial reduction to pooling. This does not mean that every detail is preserved.')

doc.add_page_break()
h('Why we chose the remaining values')
point('32 then 64 then 64 filters.', 'The first layer works on larger 28 x 28 maps. Using 32 filters instead of 64 there reduces computation and memory, with the aim of saving time. After pooling, the maps are smaller: 14 x 14 and 7 x 7. We use 64 filters in those layers to learn more feature combinations. Actual speed also depends on the number of input channels.')
point('Replicate padding.', 'We copy the nearest edge value into a one-pixel border. This lets the filter examine neighborhoods centered on the image boundary without inserting a zero border. It keeps the spatial size after convolution. The copied values are an assumption about the boundary, not new image information.')
point('GELU activation.', 'We chose GELU because it is smooth and can pass gradients for negative inputs. ReLU has a zero gradient for negative inputs. GELU may help learning in this region, but it does not guarantee better results. ReLU also supports backpropagation, and GELU gradients can still become very small.')
point('2 x 2 max pooling.', 'We keep the largest response in each small neighborhood. Pooling with stride 2 reduces 28 to 14 and then 14 to 7. We stop pooling at 7 x 7 to keep more spatial detail before classification. Max pooling keeps the strongest response, which is not always the most useful information.')
point('Dense 128 and dropout 0.3.', 'The dense layer combines the 3,136 flattened features. We retained 128 hidden units as a starting size and use GELU here too. Dropout switches off 30% of these activations during training to reduce reliance on particular features. These settings were not compared with alternatives in this run.')
point('Ten softmax outputs.', 'There are ten clothing classes, so the final layer produces ten probabilities. We use categorical cross-entropy because the labels are one-hot encoded. Lower loss means that the model assigns more probability to the correct answers.')
h('Training choices')
point('Batch size 128.', 'Assignment 1 used Iris, which has only 150 examples in total. Here we train on 54,000 images. Increasing the batch size from 32 to 128 reduces weight updates from 1,688 to 422 per epoch. Each update averages more examples, aiming for less noisy gradients and less update overhead. Smoother curves and better accuracy are not guaranteed.')
point('Forty epochs.', 'We kept the epoch count from Assignment 1. This lets us see both the early improvement and what happens when training continues for longer. The full run uses all 40 epochs.')
point('Adam with learning rate 0.001.', 'We retained these as starting training settings. Adam adjusts updates using information from earlier gradients. We did not compare different optimizers or learning rates in this experiment.')
p('The model with the lowest validation loss is kept in memory. Test values are recorded after every epoch for the graphs, but they do not update the weights or choose the best model. Seed 42 is also used for training repeatability.')

doc.add_page_break()
h('Learning curves and overfitting')
fig('learning_curves.png',6.85,'Figure 1. Loss and accuracy for training, validation and test data across all 40 epochs.')
p('During roughly the first six epochs, all three losses fall. This means that the model improves on both the training data and unseen images. Around epoch 6, validation loss is near its lowest point, at about 0.21, while test loss is about 0.25. These values are approximate readings from the graph.')
p('After this point, training loss continues to fall, but validation and test loss generally rise, with some short fluctuations. By epoch 40, training loss is about 0.03, validation loss about 0.46 and test loss about 0.53. The model fits the training examples more closely while its probability predictions on unseen images get worse. This is evidence of overfitting from around epoch 6.')
h('Why accuracy does not show this as clearly')
p('Accuracy only checks whether the class with the highest probability is correct. Cross-entropy also checks how much probability the model gives the correct class. For one image, the loss is minus the natural logarithm of that probability.')
p('For example, suppose a shirt is wrongly predicted as a coat. If the probability of the correct shirt class falls from 0.20 to 0.02, both predictions are still wrong, so accuracy does not change. However, the loss rises from about 1.61 to 3.91. A few very confident mistakes can raise the average loss while accuracy stays similar.')
p('Validation accuracy stays around 92 to 93 percent and test accuracy around 91 to 92 percent for much of the run. They do not show the same clear decline as the loss curves. However, training accuracy rises towards 99 percent, so the growing accuracy gap is also a warning sign. Loss gives the clearer signal here.')

doc.add_page_break()
h('Test predictions by clothing class')
fig('confusion_matrix.png',6.2,'Figure 2. Test confusion matrix for the weights selected using the lowest validation loss. Rows are true classes and columns are predictions.')
p('The diagonal contains correct predictions. Adding its values gives 9,113 correct predictions out of 10,000 test images, or 91.13 percent accuracy. This matrix describes the selected model, not necessarily the model at epoch 40.')
p('Bags, sneakers, sandals and trousers are classified well: 990, 987, 986 and 983 correct predictions out of 1,000 for each class. Shirts are harder, with 730 correct predictions. The model predicts 133 shirts as T-shirts or tops, and 95 pullovers as shirts. Similar clothing shapes may help explain these errors.')
h('Conclusion')
p('The CNN learns useful patterns, but training for longer does not keep improving its predictions on unseen data. The clearest warning appears in the loss curves after about epoch 6. Accuracy alone would make this problem easier to miss. For this run, selecting the weights with the lowest validation loss is more useful than simply keeping the last epoch.')
p('These results come from one run. They support the overfitting explanation, but they do not prove that each design choice is better than other options. The next exercise can compare this CNN with a model that adds residual connections.')

doc.add_page_break()
h('Exercise 2 Residual connections')
p('Exercise 2 uses a ResNet-50 model. It keeps the Fashion-MNIST data split and training settings from Exercise 1, but replaces the small three-convolution CNN with a much deeper model that contains residual connections.')
h('Part i Residual connections and gradients')
p('A residual connection sends the input around the convolutional path and adds it to the output. A bottleneck block calculates output = F(input) + shortcut(input). F is the sequence of convolutional layers. The shortcut gives gradients a more direct path to earlier layers during backpropagation. This can make a deep network easier to train, but it does not guarantee better accuracy or prevent overfitting.')
h('Part ii ResNet 50 model')
p('As shown in the lesson, ResNet-50 uses bottleneck blocks with a 1 x 1 convolution, then a 3 x 3 convolution, then another 1 x 1 convolution. We use 3, 4, 6 and 3 blocks in the four stages. This gives 16 bottleneck blocks and 49 convolutional layers before the final classifier.')
p('The first 1 x 1 convolution reduces the number of channels inside a block. The 3 x 3 convolution learns local image patterns. The final 1 x 1 convolution expands the channels before adding the shortcut. We use only ResNet-50 because ResNet-101 and ResNet-152 have many more blocks. They would take much longer to train, while the 28 x 28 Fashion-MNIST images and this exercise need one clear comparison with Exercise 1.')
table(['Part','Output size'],[
('Input','28 x 28 x 1'),
('Small 3 x 3 stem','28 x 28 x 16'),
('Stage 1: 3 blocks','28 x 28 x 64'),
('Stage 2: 4 blocks','14 x 14 x 128'),
('Stage 3: 6 blocks','7 x 7 x 256'),
('Stage 4: 3 blocks','4 x 4 x 512'),
('Global average pooling','512 values'),
('Classifier','10 probabilities')],[2.3,4.5])
p('The convolutions use zero padding at the image borders. When a shortcut needs more channels, we append channels filled with zeros. At the start of later stages, average pooling reduces the shortcut height and width before the zero channels are added. This lets the shortcut and the convolutional path have the same shape for addition.')

doc.add_page_break()
h('Part iii Training results and comparison')
point('Training time.', 'The ResNet-50 run took 12,647.5 seconds. This is much more time than the Exercise 1 CNN run because ResNet-50 has 16 bottleneck blocks and 49 convolutional layers, while Exercise 1 has only three convolutional layers.')
point('Overfitting.', 'For this ResNet run, we cannot identify one clear epoch where overfitting starts. The validation and test loss and accuracy have many rises and falls instead of changing smoothly. These fluctuations make the curve harder to interpret than the Exercise 1 curve. We should therefore not make a strong conclusion about the exact beginning of overfitting from this single run.')
point('Comparison with Exercise 1.', 'In Exercise 1, the training loss continued to decrease after about epoch 6 while validation and test loss generally increased. That gave a clearer overfitting pattern. In the ResNet graph, training loss decreases, but validation and test results move up and down across many epochs. The deeper model may need further experiments, for example changing the learning rate or repeating the run, before we can make a stronger conclusion.')
fig('learning_curves.png',6.85,'Figure 3. ResNet-50 loss and accuracy for training, validation and test data across all 40 epochs.',resnet_plots)
h('Why the ResNet curves are difficult to read')
p('The training curves improve steadily, but the validation and test curves are not smooth. For example, validation and test loss rise and fall many times after the first few epochs, while validation and test accuracy stay around a similar level with repeated changes. This means that a single loss increase is not enough to identify a clear overfitting point. The selected weights still come from the epoch with the lowest validation loss.')

doc.add_page_break()
h('ResNet test predictions by clothing class')
fig('confusion_matrix.png',6.2,'Figure 4. ResNet-50 test confusion matrix for the weights selected using the lowest validation loss. Rows are true classes and columns are predictions.',resnet_plots)
p('This confusion matrix shows which clothing classes the selected ResNet-50 model classifies correctly and which classes it confuses. It should be compared with the Exercise 1 confusion matrix together with the learning curves, rather than using training time alone to decide which model is better.')
doc.add_page_break()
h('Conclusion')
p('Exercise 1 trains much faster and shows a clear overfitting pattern after around epoch 6. ResNet-50 uses residual connections and a much deeper bottleneck structure, but it needs 12,647.5 seconds for this run. Its validation and test curves have many fluctuations, so this single run does not show a clear point where overfitting begins. The comparison shows that a deeper model increases training cost and does not automatically give smoother results.')
out=root/'assignment-2/Exercise_1_Report.docx'
doc.save(out)
print(out)
